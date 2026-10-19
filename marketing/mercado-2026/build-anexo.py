# Anexo de datos y cálculos (Excel). Uso: python3 build-anexo.py [salida.xlsx]
import json, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
D = os.path.dirname(os.path.abspath(__file__))
O = json.load(open(f"{D}/datos/mercado.json"))
INS = json.loads(subprocess.check_output(["node", "-e", f"console.log(JSON.stringify(require('{D}/insights.js')))"]))
MES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
K = {"us": "Share unidades %", "vs": "Share valor %", "pi": "Índice de precio"}
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="0B2545")
wb = Workbook()

def hoja(t, cols, filas, anchos=None):
    ws = wb.create_sheet(t[:31])
    ws.append(cols)
    for c in ws[1]: c.font, c.fill = H, HF
    for f in filas: ws.append(["Sin datos" if v is None else v for v in f])
    for i, w in enumerate(anchos or [18] * len(cols)): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "B2"
    return ws

ws = wb.active; ws.title = "Léame"
for l in ["Market Intelligence & Competitive Performance 2026 — Drean · Anexo de datos y cálculos",
          f"Fuente: GfK Brands Timeseries cargado en mercado_share; extracto del {O['lectura']}. Último mes: {O['ultimo_mes']}.",
          "MAT = acumulado móvil 12 meses (mes = último mes de la ventana). 'Sin datos' = dato no disponible en la base.",
          "Δ en puntos porcentuales (pp) para shares y en puntos para el índice de precio. No se promedian shares.",
          "No hay volúmenes, facturación ni precio absoluto del mercado: el share global y el índice ponderado no son calculables (hallazgo 18).",
          f"Validación: índice de precio = share valor / share unidades × 100 (mediana de la diferencia {O['validacion']['pi_identidad_mediana']}, p90 |dif| {O['validacion']['pi_identidad_p90_abs']:.2f}, n={O['validacion']['pi_identidad_n']}).",
          "Hojas: Matriz insights · Drean resumen · por categoría: Ranking, Mensual, MAT, Segmentos."]:
    ws.append([l])
ws.column_dimensions["A"].width = 140

w = hoja("Matriz insights", ["#", "Mercado", "Hallazgo", "Evidencia", "Comparación", "Interpretación", "Implicancia", "Tipo", "Confianza"],
         [[x["id"], x["mercado"], x["hallazgo"], x["evidencia"], x["comparacion"], x["interpretacion"], x["implicancia"], x["tipo"], x["confianza"]] for x in INS],
         [5, 13, 45, 60, 25, 50, 50, 14, 10])
for row in w.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

COLS = ["MAT jul-25", "MAT dic-25", "MAT jul-26", "Δ MAT vs jul-25", "Δ MAT vs dic-25", "Jul-25 (mes)", "Dic-25 (mes)", "Jul-26 (mes)", "Δ vs jul-25", "Δ vs dic-25"]
fila = lambda t: [t["mat25"], t["matdic"], t["mat26"], t["d_mat"], t["d_matdic"], t["jul25"], t["dic25"], t["jul26"], t["d_jul25"], t["d_dic25"]]
hoja("Drean resumen", ["Categoría", "Métrica"] + COLS,
     [[cat, K[k]] + fila(c["tabla"]["DREAN"][k]) for cat, c in O["categorias"].items() for k in K], [14, 18] + [13] * 10)
for cat, c in O["categorias"].items():
    ab = cat[:5]
    hoja(f"{ab} Ranking", ["Marca", "Métrica"] + COLS, [[b, K[k]] + fila(c["tabla"][b][k]) for b in c["marcas"] for k in K], [26, 18] + [13] * 10)
    hoja(f"{ab} Mensual", ["Marca", "Métrica", "Año"] + MES,
         [[b, K[k], y] + c["mensual"][b][k][y] for b in c["marcas"] for k in K for y in ("2025", "2026")], [26, 18, 6] + [8] * 12)
    hoja(f"{ab} MAT", ["Marca", "Métrica"] + [m[:7] for m in c["mat_meses"]],
         [[b, K[k]] + c["mat"][b][k] for b in c["marcas"] for k in K], [26, 18] + [9] * len(c["mat_meses"]))
    hoja(f"{ab} Segmentos", ["Segmento", "Marca", "Métrica"] + COLS,
         [[sg, b, K[k]] + fila(z["tabla"][b][k]) for sg, z in c["segmentos"].items() for b in z["marcas"] for k in K], [10, 26, 18] + [13] * 10)
out = sys.argv[1] if len(sys.argv) > 1 else f"{D}/anexo-market-intelligence-2026.xlsx"
wb.save(out); print("ok", out, len(wb.sheetnames), "hojas")
