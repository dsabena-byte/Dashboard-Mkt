"""Anexo analítico del diagnóstico 2026: todas las series y tablas de datos/*.json + matriz de insights + limitaciones.
Uso: python3 build-anexo.py salida.xlsx  (lee la matriz desde matriz.js vía node)."""
import json, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
MESES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct (parcial)', 'Nov', 'Dic']
NEGRITA, CAB = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='0B2545')
M = json.loads(subprocess.check_output(['node', '-e', "console.log(JSON.stringify(require('./matriz')))"], cwd=AQUI))

def hoja(wb, nombre, cols, filas, anchos=None):
    ws = wb.create_sheet(nombre[:31])
    ws.append(cols)
    for c in ws[1]: c.font, c.fill, c.alignment = NEGRITA, CAB, Alignment(wrap_text=True, vertical='top')
    for f in filas: ws.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in f])
    for i, w in enumerate(anchos or [18] * len(cols), 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws

wb = Workbook(); wb.remove(wb.active)
hoja(wb, 'Matriz de insights', ['Hallazgo', 'Evidencia', 'Comparación', 'Interpretación', 'Implicancia', 'Confianza'],
     [[m['hallazgo'], m['evidencia'], m['comparacion'], m['interpretacion'], m['implicancia'], m['confianza']] for m in M['matriz']], [34, 50, 30, 36, 36, 12])
hoja(wb, 'Limitaciones', ['Área', 'Limitación', 'Efecto en el diagnóstico', 'Cómo resolverlo'], M['limitaciones'], [14, 45, 55, 45])
hoja(wb, 'Top 10 hallazgos', ['#', 'Hallazgo', 'Tipo', 'Confianza'], [[i + 1, h['titulo'], h['tipo'], h['confianza']] for i, h in enumerate(M['top10'])], [5, 110, 18, 14])

indice = []
for area in ['social', 'pauta', 'web']:
    d = json.load(open(os.path.join(AQUI, 'datos', area + '.json')))
    filas = []
    for s in d['series']:
        for anio, k in ((2026, 'meses2026'), (2025, 'meses2025')):
            if s.get(k): filas.append([s['id'], anio, s.get('titulo_sugerido'), s.get('unidad')] + s[k] + [s.get('fuente'), s.get('notas')])
    hoja(wb, f'{area} · series', ['id', 'año', 'título', 'unidad'] + MESES + ['fuente', 'notas'], filas, [28, 6, 40, 12] + [11] * 12 + [40, 60])
    for t in d['tablas']:
        nombre = f"{area[:3]}·{t['id']}"[:31]
        hoja(wb, nombre, t['columnas'], t['filas'])
        indice.append([area, t['id'], nombre, t.get('fuente'), len(t['filas'])])
sm = json.load(open(os.path.join(AQUI, 'datos', 'sentimiento-mensual.json')))
filas = [[k, i + 1] + list(r.values()) for k, v in sm.items() for i, r in enumerate(v) if r['posts']]
hoja(wb, 'social · sentimiento mensual', ['red_año', 'mes'] + list(next(iter(sm.values()))[0].keys()), filas)
hoja(wb, 'Índice de tablas', ['área', 'id', 'hoja', 'fuente', 'filas'], indice, [10, 32, 32, 70, 8])
wb.move_sheet('Índice de tablas', offset=-(len(wb.sheetnames) - 3))
wb.save(sys.argv[1] if len(sys.argv) > 1 else 'anexo-analitico-2026.xlsx')
print('ok', len(wb.sheetnames), 'hojas')
